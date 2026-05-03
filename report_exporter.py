from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from attendance_store import AttendanceRecord, format_minutes


HEADERS = [
    "N",
    "Сотрудник",
    "Username",
    "Дата",
    "Приход",
    "Уход",
    "Отработано",
    "Минуты",
    "Статус",
    "Гео приход",
    "Гео уход",
]


def export_daily_report(path: str | Path, report_date: str, records: list[AttendanceRecord]) -> Path:
    output_path = Path(path)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Attendance"

    worksheet.append([f"Отчет посещаемости за {report_date}"])
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    title_cell = worksheet.cell(row=1, column=1)
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center")

    summary = build_summary(records)
    worksheet.append(["Сотрудников", summary["employees"]])
    worksheet.append(["На работе сейчас", summary["inside"]])
    worksheet.append(["Без прихода", summary["absent"]])
    worksheet.append(["Итого закрытых часов", format_minutes(summary["completed_minutes"])])
    worksheet.append([])
    worksheet.append(HEADERS)

    for row_index in range(2, 6):
        worksheet.cell(row=row_index, column=1).font = Font(bold=True)

    header_row = 7
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in worksheet[header_row]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for index, record in enumerate(records, start=1):
        worksheet.append(
            [
                index,
                record.full_name,
                f"@{record.username}" if record.username else "",
                record.work_date,
                record.check_in or "",
                record.check_out or "",
                format_minutes(record.worked_minutes),
                record.worked_minutes or 0,
                record_status(record),
                format_distance(record.check_in_distance_m),
                format_distance(record.check_out_distance_m),
            ]
        )

    for column_index, width in enumerate([6, 28, 20, 14, 12, 12, 18, 10, 16, 14, 14], start=1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    for row in worksheet.iter_rows(min_row=8):
        for cell in row:
            cell.alignment = Alignment(vertical="center")

    worksheet.freeze_panes = "A8"
    worksheet.auto_filter.ref = f"A7:{get_column_letter(len(HEADERS))}{max(7, worksheet.max_row)}"
    workbook.save(output_path)
    return output_path


def build_summary(records: list[AttendanceRecord]) -> dict[str, int]:
    return {
        "employees": len(records),
        "inside": sum(1 for record in records if record.is_inside),
        "absent": sum(1 for record in records if not record.check_in),
        "completed_minutes": sum(record.worked_minutes or 0 for record in records),
    }


def record_status(record: AttendanceRecord) -> str:
    if record.is_inside:
        return "на работе"
    if record.check_out:
        return "закрыто"
    return "нет прихода"


def format_distance(distance_m: int | None) -> str:
    return f"{distance_m} м" if distance_m is not None else "-"
